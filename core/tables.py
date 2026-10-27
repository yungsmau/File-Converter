import codecs
import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from core.errors import ConversionError
from core.utils import resolve_output, safe_name

Options = dict[str, Any]


@dataclass
class Sheet:
    """Промежуточный формат: любой читатель выдает список таких листов, любой писатель их принимает"""

    name: str
    rows: list[list[Any]]
    # True, если значения - необработанные строки из текстового формата

    from_text: bool = False
    decimal_comma: bool = False


# --- вспомогательное ---
def _is_empty(value: Any) -> bool:
    return value is None or value == ""


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    """Отрезает пустые строки снизу и пустые столбцы справа, выравнивает ширину строк"""
    while rows and all(_is_empty(c) for c in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for i in range(len(row) - 1, -1, -1):
            if not _is_empty(row[i]):
                width = max(width, i + 1)
                break

    return [row[:width] + [None] * (width - len(row)) for row in rows]


# --- читатели ---
def read_xlsx(path: Path, options: Options) -> list[Sheet]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ConversionError(f"Не удалось открыть Excel-файл: {e}") from e

    try:
        sheets = []
        for ws in wb.worksheets:
            # в read_only режиме размеры листа берутся из метаданных файла, а они бывает неверные
            ws.reset_dimensions()
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            sheets.append(Sheet(ws.title, _trim(rows)))
    finally:
        wb.close()
    return sheets


def read_xls(path: Path, options: Options) -> list[Sheet]:
    import xlrd

    try:
        book = xlrd.open_workbook(path)
    except Exception as e:
        raise ConversionError(f"Не удалось открыть файл формата .xls: {e}") from e

    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                t = cell.ctype
                if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    value = None
                elif t == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                elif t == xlrd.XL_CELL_BOOLEAN:
                    value = bool(cell.value)
                elif t == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(cell.value, "#ERROR")
                elif t == xlrd.XL_CELL_NUMBER:
                    # xlrd отдает все чилса как float
                    value = int(cell.value) if cell.value.is_integer() else cell.value
                else:
                    value = cell.value
                row.append(value)
            rows.append(row)
        sheets.append(Sheet(sh.name, _trim(rows)))
    return sheets


def _decode(raw: bytes) -> str:
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ConversionError("Не удалось определить кодировку файла (не utf-8, не ср1251)")


def _detect_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def read_csv(path: Path, options: Options) -> list[Sheet]:
    text = _decode(path.read_bytes())
    delimiter = _detect_delimiter(text, path.suffix.lower())
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    sheet = Sheet(
        name=path.name, rows=_trim(rows), from_text=True, decimal_comma=delimiter != ","
    )
    return [sheet]


# --- писатели ---
_INT = re.compile(r"^-?(0|[1-9]\d*)$")
_FLOAT = re.compile(r"^-?(0|[1-9]\d*)?[.,]\d+$")


def _infer(value: str, decimal_comma: bool) -> Any:
    """Строка из CSV -> чило, если это безопасно. Иначе остается строкой"""
    s = value.strip()
    if not s:
        return None
    if _INT.match(s) and len(s.lstrip("-")) <= 15:
        return int(s)
    if _FLOAT.match(s):
        if "," in s:
            return float(s.replace(",", ".")) if decimal_comma else value
        return float(s)
    return value


def _sheet_title(name: str, used: set[str]) -> str:
    """Excel: заголовок до 31 символа, без []:*?/\\, уникальный без учета регистра"""
    base = re.sub(r"[\[\]:*?/\\]", "_", name).strip("'")[:31] or "Sheet"
    title, n = base, 1
    while title.lower() in used:
        n += 1
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
    used.add(title.lower())
    return title


def write_xlsx(sheets: list[Sheet], dst: Path, options: Options) -> list[Path]:
    wb = Workbook()
    wb.remove(wb.active)  # убираем лист "Sheet", который workbook создает сам
    used: set[str] = set()

    for sheet in sheets:
        ws = wb.create_sheet(_sheet_title(sheet.name, used))
        for r, row in enumerate(sheet.rows, start=1):
            for c, value in enumerate(row, start=1):
                if sheet.from_text and isinstance(value, str):
                    value = _infer(value, sheet.decimal_comma)
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                if value is None or value == "":
                    continue
                cell = ws.cell(r, c, value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"

    target = resolve_output(dst, options.get("overwrite", False))
    wb.save(target)
    return [target]


def _csv_cell(value: Any, decimal_comma: bool) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time() == time(0)
            else value.isoformat(sep=" ")
        )
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and decimal_comma:
        return repr(value).replace(".", ",")
    return value


def write_csv(sheets: list[Sheet], dst: Path, options: Options) -> list[Path]:
    """Один лист - один файл dst. Несколько листов - по файлу на лист"""
    is_tsv = dst.suffix.lower() == ".tsv"
    delimiter = "\t" if is_tsv else options.get("delimiter", ",")
    decimal_comma = delimiter == ";"
    overwrite = options.get("overwrite", False)

    non_empty = [s for s in sheets if s.rows]
    sheets = non_empty or sheets[:1]

    outputs = []
    for sheet in sheets:
        if len(sheets) == 1:
            target = dst
        else:
            target = dst.with_name(f"{dst.stem}_{safe_name(sheet.name)}{dst.suffix}")
        target = resolve_output(target, overwrite)

        # utf-8-sig: DOM помогает Excel правильно открыть кириллицу
        with open(target, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimmiter=delimiter)
            for row in sheets.rows:
                writer.writerow(_csv_cell(v, decimal_comma) for v in row)
        outputs.append(target)
    return outputs
