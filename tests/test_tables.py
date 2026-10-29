import csv

import openpyxl
import pytest

from core.errors import ConversionError
from core.registry import convert


def read_csv(path, delimiter=","):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f, delimiter=delimiter))


def test_xlsx_to_csv(tmp_path, make_xlsx):
    src = make_xlsx(
        "t.xlsx", {"Лист1": [["имя", "сумма"], ["Аня", 10], ["Борис", 2.5]]}
    )
    [out] = convert(src, tmp_path / "t.csv", {"delimiter": ","})
    assert read_csv(out) == [["имя", "сумма"], ["Аня", "10"], ["Борис", "2.5"]]


def test_semicolon_means_decimal_comma(tmp_path, make_xlsx):
    src = make_xlsx("t.xlsx", {"Лист1": [["x"], [1.5]]})
    [out] = convert(src, tmp_path / "t.csv", {"delimiter": ";"})
    assert read_csv(out, ";")[1] == ["1,5"]


def test_several_sheets_become_several_files(tmp_path, make_xlsx):
    src = make_xlsx("t.xlsx", {"Продажи": [["а"], [1]], "Итоги": [["b"], [2]]})
    outs = convert(src, tmp_path / "t.csv", {})
    assert sorted(p.name for p in outs) == ["t_Итоги.csv", "t_Продажи.csv"]


def test_csv_to_xlsx_infers_numbers_and_keeps_leading_zeros(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text("код;сумма\n007;1,5\n123;2\n", encoding="utf-8-sig")
    [out] = convert(src, tmp_path / "out.xlsx", {})
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == "007"
    assert ws["B2"].value == 1.5
    assert ws["A3"].value == 123


def test_csv_in_cp1251_is_read(tmp_path):
    src = tmp_path / "old.csv"
    src.write_text("город;население\Москва;13\n", encoding="cp1251")
    [out] = convert(src, tmp_path / "new.tsv", {})
    assert "Москва" in out.read_text(encoding="utf-8-sig")


def test_formulas_stay_text(tmp_path):
    src = tmp_path / "f.csv"
    src.write_text("a\n=1+1\n", encoding="utf-8")
    [out] = convert(src, tmp_path / "f.xlsx", {})
    cell = openpyxl.load_workbook(out).active["A2"]
    assert cell.data_type == "s"


def test_existing_file_is_not_overwritten(tmp_path, make_xlsx):
    src = make_xlsx("t.xlsx", {"Лист1": [["a"], [1]]})
    (tmp_path / "t.csv").write_text("не трогать", encoding="utf-8")
    [out] = convert(src, tmp_path / "t.csv", {})
    assert out.name == "t (1).csv"
    assert (tmp_path / "t.csv").read_text() == "не трогать"


def test_missing_file_gives_readable_error(tmp_path):
    with pytest.raises(ConversionError, match="не найден"):
        convert(tmp_path / "нет.xlsx", tmp_path / "нет.csv", {})
